import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import cv2
import numpy as np
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import create_app
from cloud import CloudBridge
from engine import DEFAULTS, Engine
from exports import excel, report
from store import Store, now


class FakeDetector:
    def __init__(self, name="haar", settings=None, faces=True):
        self.faces = faces

    def detect(self, image):
        return [{"box": [10, 10, 40, 40], "score": None, "raw": None}] if self.faces else []


class FakeIdentity:
    def __init__(self, user_id):
        self.user_id = user_id

    def identify(self, image, face):
        return self.user_id, 30.0


class PlatformTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app(self.temp.name)
        self.client = self.app.test_client()
        self.store = self.app.extensions["store"]
        self.engine = self.app.extensions["engine"]
        response = self.client.post("/login", json={"password": self.app.config["ADMIN_PASSWORD"]})
        self.assertEqual(response.status_code, 200)
        self.token = self.client.get("/api/state").json["csrf"]
        self.headers = {"X-CSRF-Token": self.token}
        self.image = np.full((100, 100, 3), 150, np.uint8)
        self.content = cv2.imencode(".jpg", self.image)[1].tobytes()
        self.person = {"id": "1001", "name": "First Student", "class_name": "CE492"}

    def tearDown(self):
        self.engine.stop()
        self.engine.stop_camera_preview()
        self.temp.cleanup()

    def add_user(self):
        with patch("engine.Detector", FakeDetector):
            self.engine.add_user(self.person, [self.content])

    def create_session(self):
        self.store.execute("INSERT INTO sessions VALUES (?,?,?,?,NULL,'running',NULL,?)",
                           ("test", "Class session", "CE492", now(), json.dumps(DEFAULTS)))

    def test_admin_and_csrf(self):
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/api/state").status_code, 401)
        self.assertEqual(anonymous.get("/frame.jpg").status_code, 401)
        self.assertEqual(self.client.post("/api/settings", json={}).status_code, 403)
        self.assertEqual(self.client.post("/api/settings", headers=self.headers, json={"haar_scale": 1}).status_code, 400)

    def test_enrollment_appends_and_soft_removal_retains_history(self):
        self.add_user()
        self.add_user()
        self.assertEqual(self.store.users()[0]["photos"], 2)
        self.create_session()
        self.engine.process_frame("test", 1, self.image, DEFAULTS, FakeDetector(), FakeIdentity("1001"), {"1001": self.person})
        self.engine.remove_user("1001")
        self.assertFalse(self.store.users())
        self.assertEqual(self.store.query("SELECT COUNT(*) AS count FROM photos")[0]["count"], 0)
        self.assertEqual(self.store.attendance("test")[0]["name"], "First Student")

    def test_each_frame_and_match_persist_but_attendance_deduplicates(self):
        self.add_user()
        self.create_session()
        for frame_number in (1, 2):
            self.engine.process_frame("test", frame_number, self.image, DEFAULTS, FakeDetector(),
                                      FakeIdentity("1001"), {"1001": self.person})
        self.engine.process_frame("test", 3, self.image, DEFAULTS, FakeDetector(), FakeIdentity(None), {})
        self.engine.process_frame("test", 4, self.image, DEFAULTS, FakeDetector(faces=False), FakeIdentity(None), {})
        self.assertEqual(len(self.store.attendance("test")), 1)
        self.assertEqual(self.store.statistics("test")["frames"], 4)
        events = self.store.query("SELECT * FROM events WHERE session_id='test'")
        self.assertEqual(len(events), 3)
        self.assertEqual(events[-1]["accepted"], 0)
        self.assertEqual(events[-1]["user_name"], "Unknown")
        self.assertIsNone(events[0]["detector_score"])
        self.assertTrue(events[0]["image"].startswith(b"\xff\xd8"))
        self.assertTrue(events[0]["thumbnail"].startswith(b"\xff\xd8"))
        frames = self.store.query("SELECT * FROM frames")
        self.assertEqual(frames[-1]["faces"], 0)
        self.assertGreater(frames[0]["write_ms"], 0)
        self.assertIn("haar_scale", json.loads(frames[0]["config"]))

    def test_gate_logs_misses_without_calling_selected_detector(self):
        self.create_session()
        detector = FakeDetector()
        detector.detect = lambda image: self.fail("Selected detector should not run when gate sees no faces")
        self.engine.process_frame("test", 1, self.image, DEFAULTS, detector, FakeIdentity(None), {}, FakeDetector(faces=False))
        self.assertEqual(self.store.statistics("test")["frames"], 1)
        self.assertEqual(self.store.statistics("test")["faces"], 0)

    def test_recognition_display_has_one_first_detection_per_person(self):
        self.add_user()
        self.create_session()
        people = {"1001": self.person, "1002": {**self.person, "id": "1002"}}
        for frame_number, user_id in enumerate(("1001", "1001", None, "1002", "1002"), 1):
            self.engine.process_frame("test", frame_number, self.image, DEFAULTS, FakeDetector(),
                                      FakeIdentity(user_id), people)
        response = self.client.get("/api/sessions/test")
        self.assertEqual(response.status_code, 200)
        self.assertEqual([row["user_id"] for row in response.json["events"]], ["1001", "1002"])
        self.assertEqual([row["frame_id"] for row in response.json["events"]], [1, 4])
        self.assertEqual(len(response.json["attendance"]), 2)
        self.assertEqual(len(self.store.query("SELECT id FROM events")), 5)
        self.assertEqual(len(self.client.get("/api/sessions/test/events").json), 5)

    def test_attendance_unique_per_session_keeps_first_timestamp(self):
        self.add_user()
        self.create_session()
        people = {"1001": self.person}
        self.engine.process_frame("test", 1, self.image, DEFAULTS, FakeDetector(), FakeIdentity("1001"), people)
        original = self.store.attendance("test")[0]
        self.engine.process_frame("test", 2, self.image, DEFAULTS, FakeDetector(), FakeIdentity("1001"), people)
        self.assertEqual(self.store.attendance("test"), [original])
        self.assertEqual(len(self.store.query("SELECT id FROM outbox WHERE collection='attendance'")), 1)
        self.store.execute("INSERT INTO sessions VALUES (?,?,?,?,NULL,'running',NULL,?)",
                           ("second", "Next session", "CE492", now(), json.dumps(DEFAULTS)))
        self.engine.process_frame("second", 1, self.image, DEFAULTS, FakeDetector(), FakeIdentity("1001"), people)
        self.assertEqual(len(self.store.attendance("second")), 1)
        self.assertEqual(len(self.store.attendance("test")), 1)

    def test_enable_stop_enable_records_same_person_in_each_run(self):
        import time
        self.add_user()
        class FakeCamera:
            def __init__(camera, settings, source=None):
                pass
            def read(camera):
                return np.full((100, 100, 3), 150, np.uint8)
            def close(camera):
                pass
        available = {"detectors": {"haar": True}, "recognizers": {"lbph": True}}
        sessions = []
        with patch("engine.capabilities", return_value=available), patch("engine.Detector", FakeDetector), \
             patch("engine.IdentityModel", side_effect=lambda *args: FakeIdentity("1001")), patch("engine.Camera", FakeCamera):
            for mode in ("live", "live", "single", "single"):
                response = self.client.post("/api/start", headers=self.headers,
                                            json={"name": "Same class", "mode": mode})
                self.assertEqual(response.status_code, 200)
                session_id = response.json["session_id"]
                sessions.append(session_id)
                target_frames = 2 if mode == "live" else 1
                deadline = time.monotonic()+3
                while self.store.statistics(session_id)["frames"] < target_frames and time.monotonic() < deadline:
                    time.sleep(0.01)
                stopped = self.client.post("/api/stop", headers=self.headers, json={})
                self.assertEqual(stopped.status_code, 200)
                self.assertFalse(self.engine.running())
                self.assertIsNone(self.engine.error)
                self.assertEqual(len(self.store.attendance(session_id)), 1)
                self.assertEqual(len(self.store.recognized_people(session_id)), 1)
                if mode == "single":
                    self.assertEqual(self.store.statistics(session_id)["frames"], 1)
        self.assertEqual(len(set(sessions)), 4)
        self.assertEqual(len(self.store.query("SELECT * FROM attendance WHERE user_id='1001'")), 4)

    def test_export_embeds_images_and_preserves_metadata(self):
        self.add_user()
        self.create_session()
        self.engine.process_frame("test", 1, self.image, DEFAULTS, FakeDetector(), FakeIdentity("1001"), {"1001": self.person})
        workbook = load_workbook(excel(self.store, "test"))
        self.assertEqual(len(workbook["Attendance"]._images), 2)
        self.assertEqual(workbook["Attendance"]["B2"].value, "1001")
        self.assertEqual(workbook["Attendance"]["E2"].value, "haar")
        self.assertEqual(workbook["Frame statistics"].max_row, 2)
        self.assertIn("data:image/jpeg;base64,", report(self.store, "test"))
        response = self.client.get("/api/sessions/test/excel")
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response.headers["Content-Disposition"])

    def test_upload_rejects_bad_images_atomically(self):
        with patch("engine.Detector", FakeDetector):
            response = self.client.post("/api/users", headers=self.headers, data={**self.person,
                "photos": (io.BytesIO(b"bad data"), "portrait.jpg")})
        self.assertEqual(response.status_code, 400)
        self.assertFalse(self.store.users())

    def test_frame_exports_include_one_first_recognized_frame_per_person(self):
        import csv
        self.add_user()
        self.create_session()
        people = {"1001": self.person, "1002": {**self.person, "id": "1002", "name": "Second Student"}}
        for number, user_id in enumerate(("1001", "1001", None, "1002", "1002"), 1):
            self.engine.process_frame("test", number, self.image, DEFAULTS, FakeDetector(), FakeIdentity(user_id), people)
        workbook = load_workbook(excel(self.store, "test"))
        rows = list(workbook["Frame statistics"].values)
        self.assertEqual(len(rows), 3)
        header = rows[0]
        exported = [dict(zip(header, row)) for row in rows[1:]]
        self.assertEqual([row["user_id"] for row in exported], ["1001", "1002"])
        self.assertEqual([row["frame_number"] for row in exported], [1, 4])
        response = self.client.get("/api/sessions/test/statistics.csv")
        records = list(csv.DictReader(io.StringIO(response.data.decode())))
        self.assertEqual(len(records), 2)
        self.assertEqual([row["user_id"] for row in records], ["1001", "1002"])
        self.assertEqual(self.store.statistics("test")["frames"], 5)

    def test_all_session_export_includes_repeat_person_across_runs_and_empty_sessions(self):
        self.add_user()
        self.create_session()
        for session_id in ("second", "empty"):
            self.store.execute("INSERT INTO sessions VALUES (?,?,?,?,NULL,'running',NULL,?)",
                               (session_id, "Same class", "CE492", now(), json.dumps(DEFAULTS)))
        for session_id in ("test", "second"):
            self.engine.process_frame(session_id, 1, self.image, DEFAULTS, FakeDetector(),
                                      FakeIdentity("1001"), {"1001": self.person})
        response = self.client.get("/api/sessions/export-all.xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.data))
        sheet = workbook["Attendance"]
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["B2"].value, "1001")
        self.assertEqual(sheet["B3"].value, "1001")
        self.assertEqual({sheet["J2"].value, sheet["J3"].value}, {"test", "second"})
        self.assertEqual(len(sheet._images), 4)
        self.assertEqual(workbook["Session summary"].max_row, 4)
        self.assertEqual(workbook["Session configuration"].max_row, 4)
        self.assertEqual(workbook["Frame statistics"].max_row, 3)
        self.assertEqual(workbook["Recognition events"].max_row, 3)
        self.assertEqual(self.app.test_client().get("/api/sessions/export-all.xlsx").status_code, 401)

    def test_restarting_marks_unfinished_sessions_interrupted(self):
        self.create_session()
        restored = Store(Path(self.temp.name) / "attendance.sqlite3")
        self.assertEqual(restored.history()[0]["state"], "interrupted")

    def test_settings_frozen_during_session(self):
        self.engine.initializing = True
        with self.assertRaises(ValueError):
            self.engine.save_settings({"detector": "yunet"})
        self.engine.initializing = False

    def test_capture_reuses_preview_without_reopening_camera(self):
        self.engine.camera_preview_thread = type("Worker", (), {"is_alive": lambda worker: True})()
        self.engine.camera_preview_frame = self.image.copy()
        with patch("engine.Detector", FakeDetector), patch("engine.Camera") as camera:
            self.engine.add_user(self.person, [], captured=True)
            camera.assert_not_called()
        self.assertEqual(self.store.users()[0]["photos"], 1)
        self.engine.camera_preview_thread = None

    def test_preview_stream_requires_login(self):
        self.assertEqual(self.app.test_client().get("/api/camera/preview/stream").status_code, 401)

    def test_preview_controls_and_closes_camera(self):
        import time
        class FakeCamera:
            closed = False
            def __init__(camera, settings):
                pass
            def read(camera):
                return np.full((100, 100, 3), 150, np.uint8)
            def close(camera):
                FakeCamera.closed = True
        with patch("engine.Camera", FakeCamera):
            response = self.client.post("/api/camera/preview/start", headers=self.headers, json={})
            self.assertEqual(response.status_code, 200)
            deadline = time.monotonic()+2
            while self.engine.camera_preview_jpeg is None and time.monotonic() < deadline:
                time.sleep(0.01)
            self.assertTrue(self.engine.camera_preview_running())
            self.assertTrue(self.engine.camera_preview_jpeg.startswith(b"\xff\xd8"))
            with self.assertRaises(ValueError):
                self.engine.save_settings({"camera": "usb"})
            response = self.client.post("/api/camera/preview/stop", headers=self.headers, json={})
            self.assertEqual(response.status_code, 200)
            self.assertTrue(FakeCamera.closed)
            self.assertFalse(self.engine.camera_preview_running())

    def test_remote_command_is_not_replayed_after_response_failure(self):
        from datetime import datetime, timezone
        bridge = CloudBridge(self.engine)
        class Document:
            id = "command-test"
            fail_response = True
            def __init__(document):
                document.reference = document
            def to_dict(document):
                return {"action": "stop", "payload": {}, "created": datetime.now(timezone.utc)}
            def update(document, result):
                if document.fail_response:
                    document.fail_response = False
                    raise ConnectionError("Network disconnected after execution")
                document.result = result
        document = Document()
        with patch.object(self.engine, "stop", return_value={"stopping": False}) as stop:
            with self.assertRaises(ConnectionError):
                bridge._command(document)
            bridge._command(document)
            self.assertEqual(stop.call_count, 1)
            self.assertEqual(document.result["state"], "done")

    def test_remote_command_expiration(self):
        from datetime import datetime, timedelta, timezone
        bridge = CloudBridge(self.engine)
        class Document:
            id = "expired-test"
            def __init__(document):
                document.reference = document
            def to_dict(document):
                return {"action": "stop", "payload": {}, "created": datetime.now(timezone.utc)-timedelta(hours=1)}
            def update(document, result):
                document.result = result
        document = Document()
        with patch.object(self.engine, "stop") as stop:
            bridge._command(document)
            stop.assert_not_called()
            self.assertEqual(document.result["state"], "error")

    def test_approving_account_preserves_claims_and_limits_device(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        bridge = CloudBridge(self.engine)
        bridge.device = SimpleNamespace(id="classroom-pi")
        reference = Mock()
        snapshot = reference.get.return_value
        snapshot.exists = True
        snapshot.to_dict.return_value = {"device_id": "classroom-pi", "state": "pending", "email": "test@example.invalid"}
        bridge.client = Mock()
        bridge.client.collection.return_value.document.return_value = reference
        user = SimpleNamespace(email="test@example.invalid", disabled=False, custom_claims={"existing_role": "teacher"})
        with patch("firebase_admin.auth.get_user", return_value=user), patch("firebase_admin.auth.set_custom_user_claims") as claims:
            result = bridge._account_access("test-account", True)
            claims.assert_called_once_with("test-account", {"existing_role": "teacher", "face_track_admin": True,
                                                            "face_track_device": "classroom-pi"})
            self.assertEqual(result["state"], "approved")
        snapshot.to_dict.return_value = {"device_id": "other-device", "state": "pending", "email": "test@example.invalid"}
        with patch("firebase_admin.auth.set_custom_user_claims") as claims:
            with self.assertRaises(ValueError):
                bridge._account_access("test-account", True)
            claims.assert_not_called()

    def test_declining_account_does_not_grant_access(self):
        from types import SimpleNamespace
        from unittest.mock import Mock
        bridge = CloudBridge(self.engine)
        bridge.device = SimpleNamespace(id="classroom-pi")
        bridge.client = Mock()
        reference = bridge.client.collection.return_value.document.return_value
        reference.get.return_value.exists = True
        reference.get.return_value.to_dict.return_value = {"device_id": "classroom-pi", "state": "pending"}
        with patch("firebase_admin.auth.set_custom_user_claims") as claims:
            self.assertEqual(bridge._account_access("test-account", False)["state"], "rejected")
            claims.assert_not_called()


if __name__ == "__main__":
    unittest.main()
