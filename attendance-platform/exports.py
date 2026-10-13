import csv
import html
import io
import json

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill


def excel(store, session_id=None):
    all_sessions = session_id is None
    workbook = Workbook()
    attendance = workbook.active
    attendance.title = "Attendance"
    headers = ["Timestamp (UTC)", "ID", "Name", "Class", "Detector", "Recognizer", "Match score", "Face", "Detection image"]
    if all_sessions:
        headers.extend(["Session ID", "Session name"])
    attendance.append(headers)
    attendance_rows = (store.query("""SELECT a.*, s.name AS session_name, e.detector,e.recognizer,e.match_score
        FROM attendance a JOIN sessions s ON a.session_id=s.id JOIN events e ON a.event_id=e.id
        ORDER BY s.started,a.timestamp,a.id""") if all_sessions else store.attendance(session_id))
    streams = []
    for row_number, row in enumerate(attendance_rows, 2):
        values = [row[key] for key in ("timestamp", "user_id", "name", "class_name", "detector", "recognizer", "match_score")]
        if all_sessions:
            values.extend([None, None, row["session_id"], row["session_name"]])
        attendance.append(values)
        for cell in attendance[row_number]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        blobs = store.query("SELECT thumbnail,image FROM events WHERE id=?", (row["event_id"],))[0]
        for column, key, width in (("H", "thumbnail", 80), ("I", "image", 180)):
            stream = io.BytesIO(blobs[key])
            streams.append(stream)
            picture = Image(stream)
            ratio = min(width/picture.width, 90/picture.height)
            picture.width *= ratio
            picture.height *= ratio
            attendance.add_image(picture, f"{column}{row_number}")
        attendance.row_dimensions[row_number].height = 75
    for column in "ABCDEFG":
        attendance.column_dimensions[column].width = 24
    attendance.column_dimensions["H"].width = 15
    attendance.column_dimensions["I"].width = 30
    attendance.column_dimensions["J"].width = 36
    attendance.column_dimensions["K"].width = 30
    frames = workbook.create_sheet("Frame statistics")
    frame_sql = """SELECT a.session_id, a.user_id, a.name AS user_name, f.*
                   FROM attendance a JOIN events e ON a.event_id=e.id
                   JOIN frames f ON e.frame_id=f.id"""
    records = (store.query(frame_sql+" ORDER BY a.session_id,a.id") if all_sessions else
               store.query(frame_sql+" WHERE a.session_id=? ORDER BY a.id", (session_id,)))
    if records:
        frames.append(list(records[0]))
        for row in records:
            frames.append(list(row.values()))
    events = workbook.create_sheet("Recognition events")
    event_sql = """SELECT id,session_id,frame_id,timestamp,user_id,user_name,class_name,detector,recognizer,
                   box,detector_score,match_score,accepted,recognition_ms FROM events"""
    records = (store.query(event_sql+" ORDER BY id") if all_sessions else
               store.query(event_sql+" WHERE session_id=? ORDER BY id", (session_id,)))
    if records:
        events.append(list(records[0]))
        for row in records:
            events.append(list(row.values()))
    session = workbook.create_sheet("Session configuration")
    metadata = (store.query("SELECT * FROM sessions ORDER BY started,id") if all_sessions else
                store.query("SELECT * FROM sessions WHERE id=?", (session_id,)))
    if metadata:
        if all_sessions:
            session.append(list(metadata[0]))
            for row in metadata:
                session.append(list(row.values()))
        else:
            for key, value in metadata[0].items():
                session.append([key, value])
    if all_sessions:
        summary = workbook.create_sheet("Session summary")
        stat_keys = list(store.statistics(None))
        summary.append(["Session ID", "Session name", "Started (UTC)", "Ended (UTC)", "State", "Attendance count"]+stat_keys)
        for row in metadata:
            stats = store.statistics(row["id"])
            summary.append([row["id"], row["name"], row["started"], row["ended"], row["state"],
                            len(store.attendance(row["id"]))]+[stats[key] for key in stat_keys])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1754A0")
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def statistics_csv(store, session_id):
    buffer = io.StringIO()
    rows = store.query("""SELECT a.session_id, a.user_id, a.name AS user_name, f.*
                          FROM attendance a JOIN events e ON a.event_id=e.id
                          JOIN frames f ON e.frame_id=f.id
                          WHERE a.session_id=? ORDER BY a.id""", (session_id,))
    if rows:
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return buffer.getvalue()


def report(store, session_id):
    import base64
    escape = lambda value: html.escape(str(value))
    session = store.query("SELECT * FROM sessions WHERE id=?", (session_id,))
    if not session:
        raise ValueError("Session not found")
    stats = store.statistics(session_id)
    rows = []
    for row in store.attendance(session_id):
        image = store.query("SELECT image FROM events WHERE id=?", (row["event_id"],))[0]["image"]
        cells = "".join(f"<td>{escape(row[key])}</td>" for key in
                        ("timestamp", "user_id", "name", "class_name", "detector", "recognizer"))
        rows.append(f'<tr>{cells}<td><img width="220" src="data:image/jpeg;base64,{base64.b64encode(image).decode()}"></td></tr>')
    return ("<!doctype html><meta charset='utf-8'><title>Attendance report</title>"
            "<style>body{font-family:Arial;margin:30px}table{border-collapse:collapse;width:100%}"
            "td,th{border:1px solid #ccc;padding:8px}tr{break-inside:avoid}pre{white-space:pre-wrap}"
            "@media print{button{display:none}}</style><button onclick='window.print()'>Print / Save as PDF</button>"
            f"<h1>{escape(session[0]['name'])}</h1><p>Timestamps are UTC.</p>"
            f"<pre>{escape(json.dumps(stats, indent=2))}</pre><table><tr><th>Time</th><th>ID</th><th>Name</th>"
            "<th>Class</th><th>Detector</th><th>Recognizer</th><th>Recognition image</th></tr>"+
            "".join(rows)+f"</table><h2>Configuration</h2><pre>{escape(session[0]['config'])}</pre>")
